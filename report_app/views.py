import json
import requests
import memcache
import subprocess
import platform
from django.shortcuts import render
from django.core.paginator import Paginator
from django.db.models import Sum
from django.http import JsonResponse, HttpResponse
from django.conf import settings
from .models import LevelMetersData, Fillings, Users
from datetime import datetime, timedelta
from calendar import monthrange
from collections import defaultdict


# ---------------------- Константы FortMonitor ----------------------
FM_REPORT_URL = "http://31.173.168.107:8080/"
FM_URL_TEMPLATE = "http://31.173.168.107:8080/api/integration/v1/"
FM_LOGIN = "bassol_api"
FM_PASSWORD = "Bassl22052026"
FM_MC_SERVER = '127.0.0.1:11211'
FM_CACHE_TIMEOUT = 300   # 5 минут для кэширования данных заправок и остатков

MONTH_NAMES_RU = {
    1: 'Январь', 2: 'Февраль', 3: 'Март', 4: 'Апрель',
    5: 'Май', 6: 'Июнь', 7: 'Июль', 8: 'Август',
    9: 'Сентябрь', 10: 'Октябрь', 11: 'Ноябрь', 12: 'Декабрь'
}

VEHICLE_MAP = {
    '8': 'Hino',
    '9': 'Hitachi 1',
    '10': 'Hitachi 2',
    '24': 'Rimpull 1',
    '23': 'Rimpull 2',
    '25': 'Rimpull 3',
    '22': 'Vermeer 1',
    '15': 'Автокран XCMG',
    '16': 'Кран Zoomlion',
    '18': 'Камаз 65201',
    '21': 'Автопогрузчик FH',
    '29': 'Погрузчик Lonkin',
    '26': 'Подборщик',
    '32': 'Тепловоз ТЭМ7А',
    '31': 'Тепловоз ТЭМ9'
}

# ---------------------- Константы приложения ----------------------
SKIP_USER_IDS = {0, 1, 2, 3, 4, 5, 6, 39}
LIMIT_START_DATE = datetime(2026, 5, 5)


# ---------------------- Вспомогательные функции ----------------------
def ticks_to_datetime(ticks):
    if not ticks or ticks <= 0:
        return None
    try:
        epoch_ticks = 621355968000000000
        seconds = (ticks - epoch_ticks) / 10_000_000
        return datetime.fromtimestamp(seconds)
    except (ValueError, OverflowError, OSError):
        return None


def datetime_to_ticks(dt):
    epoch_ticks = 621355968000000000
    seconds = dt.timestamp()
    return int(seconds * 10_000_000 + epoch_ticks)


def check_kazs_online():
    """Проверяет доступность КАЗС по IP с кэшированием в memcached (30 секунд)"""
    mc = memcache.Client([FM_MC_SERVER], debug=0)
    cache_key = "kazs_status"
    cached = mc.get(cache_key)
    if cached is not None:
        return cached

    ip = settings.KAZS_IP
    param = "-n" if platform.system().lower() == "windows" else "-c"
    # 3 пакета, таймаут ожидания 5 секунд
    command = ["ping", param, "3", ip]
    try:
        result = subprocess.run(
            command,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=5
        )
        online = result.returncode == 0
    except (subprocess.TimeoutExpired, Exception):
        online = False

    mc.set(cache_key, online, time=30)
    return online

# ---------------------- FortMonitor функции с кэшированием ----------------------
def get_session_id() -> str:
    url = FM_URL_TEMPLATE + "connect"
    params = {"login": FM_LOGIN, "password": FM_PASSWORD}
    resp = requests.get(url, params=params, timeout=10)
    resp.raise_for_status()
    return resp.headers["SessionId"]


def get_session_status(session_id: str) -> bool:
    url = FM_URL_TEMPLATE + "ping"
    headers = {"SessionId": session_id}
    resp = requests.get(url, headers=headers, timeout=5)
    return resp.text.strip().lower() == "true"


def get_valid_session() -> str:
    mc = memcache.Client([FM_MC_SERVER], debug=0)
    session_id = mc.get("fortmonitor_session_id")
    if not session_id or not get_session_status(session_id):
        session_id = get_session_id()
        mc.set("fortmonitor_session_id", session_id, time=900)
    return session_id


def get_fuelings_from_fortmonitor(vehicle_id: str, date_from: datetime, date_to: datetime) -> list:
    cache_key = f"fm_fuelings_{vehicle_id}_{date_from.strftime('%Y%m%d')}_{date_to.strftime('%Y%m%d')}"
    mc = memcache.Client([FM_MC_SERVER], debug=0)
    cached = mc.get(cache_key)
    if cached is not None:
        return cached

    session_id = get_valid_session()
    url = FM_URL_TEMPLATE + "fuelings"
    headers = {"SessionId": session_id}
    from_str = date_from.strftime("%Y-%m-%d 00:00:00")
    to_str = date_to.strftime("%Y-%m-%d 23:59:59")
    params = {
        "oid": vehicle_id,
        "from": from_str,
        "to": to_str
    }
    resp = requests.get(url, headers=headers, params=params, timeout=15)
    resp.raise_for_status()
    data = resp.json()
    if data.get("result") != "Ok":
        raise Exception(f"API error: {data.get('result')}")
    fuelings = data.get("fuelings", [])
    mc.set(cache_key, fuelings, time=FM_CACHE_TIMEOUT)
    return fuelings


def get_total_fortmonitor_fuelings(vehicle_id: str, date_from: datetime, date_to: datetime) -> float:
    fuelings = get_fuelings_from_fortmonitor(vehicle_id, date_from, date_to)
    total = sum(f['volume'] for f in fuelings if f.get('fuel_type') == 'fueling')
    return total


def get_fortmonitor_fuel_level(vehicle_id: str, date_from: datetime, date_to: datetime):
    """
    Возвращает кортеж (endLevel, last_event_time).
    - endLevel: остаток топлива в баке (литры) на дату date_to (конец периода).
    - last_event_time: datetime максимального stop_time из всех событий (заправок/сливов)
      в пределах периода [date_from, date_to]. Если событий нет, last_event_time = None.

    Логика выбора датчика:
        * Сначала ищем все датчики, в названии которых есть "Датчик уровня топлива"
          (например, "Датчик уровня топлива", "Датчик уровня топлива 1", "Датчик уровня топлива 2").
          Из них берём endLevel, отбрасывая недостоверные значения (0 или 1).
          Если есть хотя бы одно валидное значение, вычисляем среднее арифметическое.
        * Если таких датчиков нет, ищем датчики, название которых начинается с "Бак"
          (например, "Бак 1", "Бак 2").
          Аналогично отбрасываем 0/1 и усредняем.
        * Если ничего не найдено – возвращаем (None, None).
    """
    cache_key = f"fm_fuel_level_{vehicle_id}_{date_from.strftime('%Y%m%d')}_{date_to.strftime('%Y%m%d')}"
    mc = memcache.Client([FM_MC_SERVER], debug=0)
    cached = mc.get(cache_key)
    if cached is not None:
        return cached

    session_id = get_valid_session()
    url = FM_URL_TEMPLATE + "getobjectsfuelinfo"
    headers = {"SessionId": session_id}
    params = {
        "objuids": vehicle_id,
        "date_from": date_from.strftime("%Y-%m-%d 00:00:00"),
        "date_to": date_to.strftime("%Y-%m-%d 23:59:59")
    }
    resp = requests.get(url, headers=headers, params=params, timeout=15)
    resp.raise_for_status()
    data = resp.json()
    if not data or not isinstance(data, list):
        return (None, None)

    obj = next((item for item in data if str(item.get('object_id')) == str(vehicle_id)), None)
    if not obj or 'sensors' not in obj:
        return (None, None)

    sensors = obj['sensors']
    end_level = None
    last_event_time = None

    # --- Сначала ищем датчики, содержащие "Датчик уровня топлива" ---
    level_sensors = [s for s in sensors if 'Датчик уровня топлива' in s.get('sensor_name', '')]
    valid_levels = []
    for s in level_sensors:
        level = s.get('endLevel')
        if level is not None and level > 1:   # игнорируем 0 и 1
            valid_levels.append(level)
    if valid_levels:
        end_level = sum(valid_levels) / len(valid_levels)
    else:
        # --- Иначе ищем датчики, начинающиеся с "Бак" ---
        tank_sensors = [s for s in sensors if s.get('sensor_name', '').startswith('Бак')]
        valid_levels = []
        for s in tank_sensors:
            level = s.get('endLevel')
            if level is not None and level > 1:
                valid_levels.append(level)
        if valid_levels:
            end_level = sum(valid_levels) / len(valid_levels)

    # Определяем время последнего события в пределах периода (из всех датчиков)
    for sensor in sensors:
        for fueling in sensor.get('fuelings', []):
            stop_time_str = fueling.get('stop_time')
            if stop_time_str:
                try:
                    dt = datetime.strptime(stop_time_str, '%Y-%m-%d %H:%M:%S')
                    if last_event_time is None or dt > last_event_time:
                        last_event_time = dt
                except:
                    pass

    result = (end_level, last_event_time)
    mc.set(cache_key, result, time=FM_CACHE_TIMEOUT)
    return result

# ---------------------- Данные по уровнемерам ----------------------
def get_fuel_balance_data():
    desired_ids = [1, 2, 3, 4]
    measurements = []
    total_liters = 0

    for lm_id in desired_ids:
        latest = LevelMetersData.objects.filter(
            id_level_meter_id=lm_id,
            fuel_volume_valid=True
        ).order_by('-date_time').first()

        if latest:
            liters = latest.fuel_volume * 1000
            total_liters += liters
            level_cm = None
            if latest.level is not None and latest.level_valid:
                level_cm = float(latest.level) * 100
            mass_kg = None
            if latest.mass is not None and latest.mass_valid:
                mass_kg = int(latest.mass * 1000)
            measurements.append({
                'id': lm_id,
                'liters': int(liters),
                'level_cm': level_cm,
                'mass_kg': mass_kg,
            })
        else:
            measurements.append({
                'id': lm_id,
                'liters': None,
                'level_cm': None,
                'mass_kg': None,
            })

    return {
        'total_volume': int(total_liters),
        'measurements': measurements,
    }


def fuel_balance_api(request):
    data = get_fuel_balance_data()
    return JsonResponse(data)

# ---------------------- Основная страница списка заправок ----------------------
def fillings_list(request):
    balance_data = get_fuel_balance_data()
    search_query = request.GET.get('search', '').strip()

    fillings = Fillings.objects.select_related(
        'id_user', 'id_controller', 'id_car', 'id_fuel'
    ).filter(litre__gt=0)

    if search_query:
        fillings = fillings.filter(id_user__full_name__icontains=search_query)

    fillings = fillings.order_by('-date_time')

    now = datetime.now()
    month_start = datetime(now.year, now.month, 1)
    ticks_month_start = datetime_to_ticks(month_start)

    for filling in fillings:
        filling.dt = ticks_to_datetime(filling.date_time)
        user = filling.id_user
        if user and user.id not in SKIP_USER_IDS:
            filling.month_limit = user.month_limit
        else:
            filling.month_limit = None

    user_ids = set()
    for filling in fillings:
        user = filling.id_user
        if (user is not None and
            user.id not in SKIP_USER_IDS and
            filling.month_limit is not None and
            filling.date_time is not None and
            filling.date_time > 0 and
            filling.dt and
            filling.dt >= LIMIT_START_DATE):
            user_ids.add(user.id)

    if user_ids:
        month_expenses = {
            item['id_user']: item['total']
            for item in Fillings.objects.filter(
                id_user__in=user_ids,
                litre__gt=0,
                date_time__gte=ticks_month_start,
                date_time__isnull=False,
                date_time__gt=0
            ).values('id_user').annotate(total=Sum('litre'))
        }
    else:
        month_expenses = {}

    for filling in fillings:
        user = filling.id_user
        filling.remaining = None
        if (user is None or
            user.id in SKIP_USER_IDS or
            filling.month_limit is None or
            filling.date_time is None or
            filling.date_time <= 0 or
            not filling.dt or
            filling.dt < LIMIT_START_DATE):
            continue

        spent = month_expenses.get(user.id, 0)
        filling.remaining = filling.month_limit - spent

    paginator = Paginator(fillings, 20)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    kazs_online = check_kazs_online()

    context = {
        'page_obj': page_obj,
        'total_volume': balance_data['total_volume'],
        'measurements': balance_data['measurements'],
        'search_query': search_query,
        'kazs_online': kazs_online,
        'low_fuel_threshold': settings.LOW_FUEL_THRESHOLD,
        'warning_threshold': settings.WARNING_THRESHOLD,
        'fuel_threshold': settings.FUEL_THRESHOLD,
    }
    return render(request, 'fillings_list.html', context)

# ---------------------- Отчёт по выдаче топлива и заправкам ----------------------
def fuel_report(request):
    users = Users.objects.filter(
        id__in=Fillings.objects.filter(litre__gt=0).values('id_user').distinct()
    ).order_by('full_name')

    selected_user_id = request.GET.get('user_id')
    date_from_str = request.GET.get('date_from')
    date_to_str = request.GET.get('date_to')
    
    total_litres = None
    selected_user = None
    remaining = None
    error = None
    fortmonitor_total = None
    is_mapped = False
    fuel_level = None
    fuel_level_time = None

    if selected_user_id and date_from_str and date_to_str:
        try:
            selected_user = Users.objects.get(id=selected_user_id)
            date_from = datetime.strptime(date_from_str, '%Y-%m-%d')
            date_to = datetime.strptime(date_to_str, '%Y-%m-%d')
            
            if date_from > date_to:
                error = "Ошибка: начальная дата не может быть позже конечной."
            else:
                date_to_end = datetime(date_to.year, date_to.month, date_to.day, 23, 59, 59)
                ticks_from = datetime_to_ticks(date_from)
                ticks_to = datetime_to_ticks(date_to_end)
                
                total = Fillings.objects.filter(
                    id_user=selected_user,
                    litre__gt=0,
                    date_time__gte=ticks_from,
                    date_time__lte=ticks_to
                ).aggregate(total=Sum('litre'))['total']
                total_litres = total if total is not None else 0

                if selected_user.id not in SKIP_USER_IDS and selected_user.month_limit is not None:
                    if date_from.year == date_to.year and date_from.month == date_to.month:
                        month_start = datetime(date_from.year, date_from.month, 1)
                        last_day = monthrange(date_from.year, date_from.month)[1]
                        month_end = datetime(date_from.year, date_from.month, last_day, 23, 59, 59)
                        ticks_month_start = datetime_to_ticks(month_start)
                        ticks_month_end = datetime_to_ticks(month_end)
                        total_month = Fillings.objects.filter(
                            id_user=selected_user,
                            litre__gt=0,
                            date_time__gte=ticks_month_start,
                            date_time__lte=ticks_month_end
                        ).aggregate(total=Sum('litre'))['total'] or 0
                        remaining = selected_user.month_limit - total_month
                    else:
                        remaining = None
                else:
                    remaining = None

                vehicle_id = None
                for vid, name in VEHICLE_MAP.items():
                    if name == selected_user.short_name:
                        vehicle_id = vid
                        break
                is_mapped = (vehicle_id is not None)
                if vehicle_id:
                    try:
                        fortmonitor_total = get_total_fortmonitor_fuelings(vehicle_id, date_from, date_to)
                        fuel_level, fuel_level_time = get_fortmonitor_fuel_level(vehicle_id, date_from, date_to)
                    except Exception as e:
                        print(f"FortMonitor error: {e}")
                        fortmonitor_total = None
                        fuel_level = None
                        fuel_level_time = None
                else:
                    fortmonitor_total = None
                    fuel_level = None
                    fuel_level_time = None

        except Users.DoesNotExist:
            error = "Выбранная машина не найдена"
        except Exception as e:
            error = f"Ошибка: {str(e)}"
    
    context = {
        'users': users,
        'selected_user_id': selected_user_id,
        'selected_user': selected_user,
        'date_from': date_from_str,
        'date_to': date_to_str,
        'total_litres': total_litres,
        'remaining': remaining,
        'error': error,
        'fortmonitor_total': fortmonitor_total,
        'is_mapped': is_mapped,
        'fuel_level': fuel_level,
        'fuel_level_time': fuel_level_time,
        'fm_report_url': FM_REPORT_URL,
        'vehicle_names': list(VEHICLE_MAP.values()),
    }
    return render(request, 'fuel_report.html', context)


def kazs_status_api(request):
    """Возвращает JSON с состоянием КАЗС для AJAX-обновления"""
    online = check_kazs_online()
    return JsonResponse({'online': online})

# ---------------------- Статистика по КАЗС ----------------------
def _collect_kazs_statistics(date_from, date_to):
    """Собирает все данные для страницы статистики по КАЗС."""
    ticks_from = datetime_to_ticks(date_from)
    ticks_to = datetime_to_ticks(date_to)

    fillings = Fillings.objects.filter(
        litre__gt=0,
        date_time__gte=ticks_from,
        date_time__lte=ticks_to,
        date_time__isnull=False,
        date_time__gt=0
    ).select_related('id_user').order_by('date_time')

    monthly = defaultdict(lambda: {'liters': 0.0, 'count': 0})
    weekday = defaultdict(float)
    hour = defaultdict(float)
    total_by_user = {}

    for f in fillings:
        dt = ticks_to_datetime(f.date_time)
        if not dt:
            continue
        key = dt.strftime('%Y-%m')
        liters = float(f.litre)
        monthly[key]['liters'] += liters
        monthly[key]['count'] += 1
        weekday[dt.weekday()] += liters
        hour[dt.hour] += liters
        if f.id_user:
            name = f.id_user.full_name
            total_by_user[name] = total_by_user.get(name, 0) + liters

    # Список месяцев выбранного периода
    chart_labels, chart_data, table_data = [], [], []
    cursor = datetime(date_from.year, date_from.month, 1)
    end_cursor = datetime(date_to.year, date_to.month, 1)
    while cursor <= end_cursor:
        key = cursor.strftime('%Y-%m')
        label = f"{MONTH_NAMES_RU[cursor.month]} {cursor.year}"
        liters = monthly[key]['liters']
        count = monthly[key]['count']
        chart_labels.append(label)
        chart_data.append(round(liters, 2))
        table_data.append({
            'month': label,
            'liters': round(liters, 2),
            'count': count,
        })
        if cursor.month == 12:
            cursor = datetime(cursor.year + 1, 1, 1)
        else:
            cursor = datetime(cursor.year, cursor.month + 1, 1)

    # Топ-15 техники
    top_users = sorted(total_by_user.items(), key=lambda x: -x[1])[:15]

    # Дни недели и часы
    weekday_labels = ['Пн', 'Вт', 'Ср', 'Чт', 'Пт', 'Сб', 'Вс']
    weekday_data = [round(weekday[i], 2) for i in range(7)]
    hour_labels = [f"{i:02d}" for i in range(24)]
    hour_data = [round(hour[i], 2) for i in range(24)]

    # Сравнение год к году
    current_year = date_to.year
    prev_year = current_year - 1
    year_agg = defaultdict(float)
    start_compare = datetime(prev_year, 1, 1)
    end_compare = datetime(current_year, 12, 31, 23, 59, 59)
    compare_fillings = Fillings.objects.filter(
        litre__gt=0,
        date_time__gte=datetime_to_ticks(start_compare),
        date_time__lte=datetime_to_ticks(end_compare),
        date_time__isnull=False,
        date_time__gt=0
    )
    for f in compare_fillings:
        dt = ticks_to_datetime(f.date_time)
        if not dt:
            continue
        year_agg[dt.strftime('%Y-%m')] += float(f.litre)

    compare_labels = [MONTH_NAMES_RU[m] for m in range(1, 13)]
    current_year_data = [round(year_agg.get(f"{current_year}-{m:02d}", 0), 2) for m in range(1, 13)]
    prev_year_data = [round(year_agg.get(f"{prev_year}-{m:02d}", 0), 2) for m in range(1, 13)]

    return {
        'chart_labels': chart_labels,
        'chart_data': chart_data,
        'table_data': table_data,
        'top_users': top_users,
        'total_period': round(sum(chart_data), 2),
        'total_fillings': sum(r['count'] for r in table_data),
        'weekday_labels': weekday_labels,
        'weekday_data': weekday_data,
        'hour_labels': hour_labels,
        'hour_data': hour_data,
        'compare_labels': compare_labels,
        'current_year': current_year,
        'prev_year': prev_year,
        'current_year_data': current_year_data,
        'prev_year_data': prev_year_data,
    }


def kazs_statistics(request):
    """Страница статистики по КАЗС с выбором периода."""
    today = datetime.now().date()
    date_from_str = request.GET.get('date_from')
    date_to_str = request.GET.get('date_to')

    if date_from_str:
        try:
            date_from = datetime.strptime(date_from_str, '%Y-%m-%d')
        except ValueError:
            date_from = datetime(today.year, 1, 1)
    else:
        date_from = datetime(today.year, 1, 1)

    if date_to_str:
        try:
            date_to_dt = datetime.strptime(date_to_str, '%Y-%m-%d')
            date_to = datetime(date_to_dt.year, date_to_dt.month, date_to_dt.day, 23, 59, 59)
        except ValueError:
            date_to = datetime(today.year, today.month, today.day, 23, 59, 59)
    else:
        date_to = datetime(today.year, today.month, today.day, 23, 59, 59)

    if date_from > date_to:
        date_from, date_to = date_to, date_from

    stats = _collect_kazs_statistics(date_from, date_to)
    stats['date_from'] = date_from.strftime('%Y-%m-%d')
    stats['date_to'] = date_to.strftime('%Y-%m-%d')
    stats['period_start'] = date_from.strftime('%d.%m.%Y')
    stats['period_end'] = date_to.strftime('%d.%m.%Y')

    return render(request, 'kazs_statistics.html', stats)


def kazs_statistics_export(request):
    """Экспорт статистики в Excel."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        return HttpResponse(
            "Библиотека openpyxl не установлена. Выполните: pip install openpyxl",
            status=500
        )

    today = datetime.now().date()
    date_from_str = request.GET.get('date_from')
    date_to_str = request.GET.get('date_to')

    if date_from_str:
        try:
            date_from = datetime.strptime(date_from_str, '%Y-%m-%d')
        except ValueError:
            date_from = datetime(today.year, 1, 1)
    else:
        date_from = datetime(today.year, 1, 1)

    if date_to_str:
        try:
            date_to_dt = datetime.strptime(date_to_str, '%Y-%m-%d')
            date_to = datetime(date_to_dt.year, date_to_dt.month, date_to_dt.day, 23, 59, 59)
        except ValueError:
            date_to = datetime(today.year, today.month, today.day, 23, 59, 59)
    else:
        date_to = datetime(today.year, today.month, today.day, 23, 59, 59)

    if date_from > date_to:
        date_from, date_to = date_to, date_from

    stats = _collect_kazs_statistics(date_from, date_to)

    wb = Workbook()

    # Лист "Помесячно"
    ws1 = wb.active
    ws1.title = "Помесячно"
    ws1.append(['Месяц', 'Литров', 'Заправок'])
    for row in stats['table_data']:
        ws1.append([row['month'], row['liters'], row['count']])

    # Лист "Топ техники"
    ws2 = wb.create_sheet("Топ техники")
    ws2.append(['№', 'Техника', 'Литров'])
    for i, (name, liters) in enumerate(stats['top_users'], start=1):
        ws2.append([i, name, round(liters, 2)])

    # Лист "Сводка"
    ws3 = wb.create_sheet("Сводка")
    ws3.append(['Период', f"{stats['period_start'] if 'period_start' in stats else date_from.strftime('%d.%m.%Y')} – {date_to.strftime('%d.%m.%Y')}"])
    ws3.append(['Всего литров', stats['total_period']])
    ws3.append(['Всего заправок', stats['total_fillings']])

    # Стили
    for ws in [ws1, ws2, ws3]:
        if ws.max_row >= 1:
            for cell in ws[1]:
                if cell.value:
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill(start_color='DDDDDD', end_color='DDDDDD', fill_type='solid')
        for column in ws.columns:
            letter = column[0].column_letter
            max_length = max(
                (len(str(c.value)) for c in column if c.value is not None),
                default=10
            )
            ws.column_dimensions[letter].width = min(max_length + 4, 50)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f"kazs_statistics_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response